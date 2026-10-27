"""Regenerate the small PPTX fixture. Requires python-pptx and Pillow; tests use
only the committed fixture, so Python is not needed to run the project tests.
"""
from pathlib import Path
from io import BytesIO
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from PIL import Image
from lxml import etree

presentation = Presentation()
presentation.slide_width = Inches(10)
presentation.slide_height = Inches(5.625)
slide = presentation.slides.add_slide(presentation.slide_layouts[5])
slide.shapes.title.text = 'Quarterly briefing 合同'
text = slide.shapes.add_textbox(Inches(.7), Inches(1.3), Inches(8), Inches(1.5))
text.text_frame.text = 'Read-only presentation preview'
paragraph = text.text_frame.add_paragraph()
paragraph.text = '<img src=x onerror="parent.previewExecuted=true">'
paragraph.font.size = Pt(15)
paragraph = text.text_frame.add_paragraph()
run = paragraph.add_run()
run.text = 'External hyperlink'
run.hyperlink.address = 'https://preview-external.invalid/link'
image = BytesIO()
Image.new('RGB', (60, 40), '#2d6547').save(image, format='PNG')
image.seek(0)
slide.shapes.add_picture(image, Inches(.7), Inches(3.2), width=Inches(1.2))
second = presentation.slides.add_slide(presentation.slide_layouts[5])
second.shapes.title.text = 'Budget summary'
table = second.shapes.add_table(3, 2, Inches(.7), Inches(1.3), Inches(3), Inches(1.8)).table
for row, values in enumerate([['Item', 'Amount'], ['Mission', '1,234.50'], ['Travel', '2,469.00']]):
    for col, value in enumerate(values):
        table.cell(row, col).text = value
chart = CategoryChartData()
chart.categories = ['One', 'Two']
chart.add_series('Amount', [12, 24])
second.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4.3), Inches(1.3), Inches(4.5), Inches(3.4), chart)
output = BytesIO()
presentation.save(output)
# Exercise the renderer's blocked external-media relationship path as well.
with ZipFile(output) as source:
    files = {name: source.read(name) for name in source.namelist()}
root = etree.fromstring(files['ppt/slides/slide1.xml'])
ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main', 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
picture = deepcopy(root.find('.//p:pic', ns))
identity = picture.find('p:nvPicPr/p:cNvPr', ns)
identity.set('id', '999')
identity.set('name', 'Blocked external image')
blip = picture.find('p:blipFill/a:blip', ns)
rel_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
blip.attrib.pop('{' + rel_ns + '}embed')
blip.set('{' + rel_ns + '}link', 'rIdExternal')
root.find('.//p:spTree', ns).append(picture)
files['ppt/slides/slide1.xml'] = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
relationships = etree.fromstring(files['ppt/slides/_rels/slide1.xml.rels'])
relationship = etree.SubElement(relationships, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship')
relationship.set('Id', 'rIdExternal')
relationship.set('Type', rel_ns + '/image')
relationship.set('Target', 'https://preview-external.invalid/tracking.png')
relationship.set('TargetMode', 'External')
files['ppt/slides/_rels/slide1.xml.rels'] = etree.tostring(relationships, xml_declaration=True, encoding='UTF-8', standalone=True)
path = Path('tests/browser/fixtures/office-preview.pptx')
with ZipFile(path, 'w', ZIP_DEFLATED) as target:
    for name, content in files.items():
        target.writestr(name, content)
print('Created', path, path.stat().st_size, 'bytes')
